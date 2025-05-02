import streamlit as st
import pandas as pd
import plotly.express as px
from pptx import Presentation
from pptx.util import Inches
import io
from datetime import datetime

# ✅ MUST BE FIRST Streamlit COMMAND
st.set_page_config(layout="wide")

# ---- Session State ----
if "selected_tile" not in st.session_state:
    st.session_state.selected_tile = None
if "page" not in st.session_state:
    st.session_state.page = None
if "objective" not in st.session_state:
    st.session_state.objective = "Objective: Not defined"

# ---- App Title ----
st.title("Business Planning Tool")

# ---- Screen 1A: Tile Selection ----
tiles = {
    "Business Overview": "CXO",
    "Financial Planning": "CFO Role",
    "Assortment Planning": "Merchandiser Role",
    "Inventory Planning": "Allocator Role",
    "Buy Planning": "Buyer Role",
    "Reorder Planning": "Demand Planner Role",
    "Markdown Planning": "Pricing Analyst Role"
}

rows = list(tiles.items())

st.markdown("### Select a Role (Only Allocator is active)")

col1, col2, col3, col4 = st.columns(4)
for idx, (title, role) in enumerate(rows[:4]):
    with [col1, col2, col3, col4][idx]:
        if title == "Inventory Planning":
            if st.button(f"**{title}**\n\n{role}"):
                st.session_state.selected_tile = title
        else:
            st.markdown(f"**{title}**\n\n_Coming Soon_")

col1, col2, col3 = st.columns(3)
for idx, (title, role) in enumerate(rows[4:]):
    with [col1, col2, col3][idx]:
        if title == "Inventory Planning":
            if st.button(f"**{title}**\n\n{role}"):
                st.session_state.selected_tile = title
        else:
            st.markdown(f"**{title}**\n\n_Coming Soon_")

# ---- Screen 1B / 1C: Use Case Selection ----
if st.session_state.get("selected_tile") == "Inventory Planning":
    st.subheader("Allocator Use Cases")
    use_cases = ["Replenishment", "Allocation", "IST", "Pullback"]
    selected = st.multiselect("Select Use Cases", use_cases)

    # Objective logic
    usecase_map = {
        ("Allocation", "Pullback", "Replenishment"): "User is doing replenishment with replacement and pullback to warehouse of bottom seller styles",
        ("Allocation", "IST", "Pullback", "Replenishment"): "User is doing ideal Inter store transfer with health correction, replenishment, replacement and pullback",
        ("Pullback", "Replenishment"): "User is doing replenishment with pullback to warehouse of bottom seller styles",
        ("Allocation", "Replenishment"): "User is doing replenishment with replacement with new styles",
    }

    usecase_tuple = tuple(sorted(selected))
    st.session_state.objective = usecase_map.get(
        usecase_tuple, f"User selected: {', '.join(selected)}"
    )
    st.markdown(f"**{st.session_state.objective}**")

    col1, col2 = st.columns(2)
    with col1:
        if st.button("➡️ Go to Dashboard View"):
            st.session_state.page = "dashboard"
    with col2:
        if st.button("📁 Go to Data Upload Page"):
            st.session_state.page = "upload"

# ---- Screen 4: File Upload ----
if st.session_state.page == "upload":
    st.subheader("📁 Upload Your Input Files (.xlsx only)")
    files = {}
    file_names = [
        "Sales", "Store Stock On Hand", "Warehouse Stock",
        "GRN", "Output", "Implementation Data", "Master"
    ]
    for name in file_names:
        files[name] = st.file_uploader(f"Upload {name} file", type=["xlsx"])

    if st.button("Process Files"):
        st.success("Files uploaded. You can now run the algorithm.")

# ---- Screen 2: Dashboard View ----
if st.session_state.page == "dashboard":
    st.subheader("📊 Dashboard: Impact Analysis")
    st.markdown(f"**{st.session_state.objective}**")

    # Inputs
    col1, col2, col3, col4 = st.columns(4)
    with col1: date = st.date_input("📅 Date")
    with col2: timestamp = st.time_input("⏱️ Time")
    with col3: dur_type = st.selectbox("Duration Type", ["Custom Range", "Fixed Days"])
    with col4:
        if dur_type == "Custom Range":
            pre_start = st.date_input("Pre Start")
            pre_end = st.date_input("Pre End")
        else:
            pre_days = st.number_input("Pre Days", min_value=1, step=1)

    # Metrics
    st.markdown("### Key Metrics (dummy data)")
    metrics = {
        "Revenue/Day": 12500,
        "ROS": 2.3,
        "Fill Rate": "86%",
        "Stock Out Rate": "14%",
        "DOH": 23,
        "Stock Health": "Good"
    }
    cols = st.columns(len(metrics))
    for i, (k, v) in enumerate(metrics.items()):
        with cols[i]:
            st.metric(k, value=v)

    # KPI Widget Table
    st.markdown("### KPI Widgets")
    kpi_df = pd.DataFrame({
        "KPI": ["rev/day", "DOH", "Stock out rate", "Sales vs Stock mix"],
        "Value": [12000, 25, "15%", "Balanced"]
    })
    st.dataframe(kpi_df)

    # Simple Bar Chart
    fig = px.bar(x=["rev/day", "ros"], y=[12000, 2.1], labels={"x": "Metric", "y": "Value"})
    st.plotly_chart(fig)

    # PPT Generator Function
    def generate_ppt():
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[0])
        title, content = slide.shapes.title, slide.placeholders[1]
        title.text = "Impact Analysis Summary"
        content.text = st.session_state.objective

        slide = ppt.slides.add_slide(ppt.slide_layouts[1])
        title, content = slide.shapes.title, slide.placeholders[1]
        title.text = "Key Metrics"
        content.text = "\n".join([f"{k}: {v}" for k, v in metrics.items()])

        buffer = io.BytesIO()
        ppt.save(buffer)
        buffer.seek(0)
        return buffer

    if st.button("📥 Download PPT"):
        ppt_file = generate_ppt()
        st.download_button("Click to Download", ppt_file, file_name="impact_summary.pptx")

